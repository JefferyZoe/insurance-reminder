#!/usr/bin/env python3
"""
生成保单详情 HTML 页面（带密码保护），部署到 GitHub Pages。
保单文件通过百度云链接查看/下载。
支持从 Excel 读取现金价值数据并汇总。
"""
import json
import os
import datetime
import hashlib
import secrets
import base64

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(SCRIPT_DIR, "insurance_data.json")
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "docs")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "index.html")
TEMPLATES_DIR = os.path.join(SCRIPT_DIR, "templates")

PAGE_PASSWORD = os.environ.get('PAGE_PASSWORD', '123456')
POLICY_FILES_DIR = os.environ.get('POLICY_FILES_DIR', os.path.join(SCRIPT_DIR, "policy_files"))
EXCEL_FILE = os.path.join(POLICY_FILES_DIR, "insurance.xlsx")


def load_config():
    with open(CONFIG_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def load_cash_values():
    """
    从 Excel 读取现金价值数据。
    Excel 结构：
    - 行2: policy_id (如 POL001, POL002, ...)
    - 行3: 类型 (guaranteed)
    - 行4+: 数据，A列为保单年度(1,2,3...)，其他列为现金价值
    返回: {policy_id: {年度: 现金价值}}
    """
    if not os.path.exists(EXCEL_FILE):
        return {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
        ws = wb.active

        # 读取第2行获取 policy_id
        row2 = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]

        # 找出每个 policy_id 对应的列索引
        policy_cols = {}  # {policy_id: col_index}
        for col_idx, val in enumerate(row2):
            if val and str(val).startswith("POL"):
                policy_cols[str(val)] = col_idx

        # 从第4行开始读取数据
        result = {}  # {policy_id: {年度: 现金价值}}
        for policy_id in policy_cols:
            result[policy_id] = {}

        for row in ws.iter_rows(min_row=4, values_only=True):
            year_val = row[0]
            if year_val is None:
                continue
            try:
                year = int(year_val)
            except (ValueError, TypeError):
                continue
            for policy_id, col_idx in policy_cols.items():
                if col_idx < len(row) and row[col_idx] is not None:
                    try:
                        result[policy_id][year] = float(row[col_idx])
                    except (ValueError, TypeError):
                        pass

        wb.close()
        return result
    except Exception:
        return {}


def calculate_policy_year(first_insure_date_str):
    """根据 first_insure_date 计算当前保单年度（first_insure_date 是第1年度）"""
    first_date = datetime.datetime.strptime(first_insure_date_str, "%Y-%m-%d").date()
    today = datetime.date.today()
    # 当前保单年度 = 今年 - 首保年份 + 1
    # 如果还没到今年的保单周年日，算上一年度
    try:
        anniversary = first_date.replace(year=today.year)
    except ValueError:
        anniversary = first_date.replace(year=today.year, day=28)
    if today >= anniversary:
        return today.year - first_date.year + 1
    else:
        return today.year - first_date.year


def load_template(filename):
    with open(os.path.join(TEMPLATES_DIR, filename), "r", encoding="utf-8") as f:
        return f.read()


def calculate_next_due_date(first_insure_date_str, pay_period_years):
    first_date = datetime.datetime.strptime(first_insure_date_str, "%Y-%m-%d").date()
    today = datetime.date.today()
    end_year = first_date.year + pay_period_years - 1
    for year in range(first_date.year + 1, end_year + 1):
        try:
            due_date = first_date.replace(year=year)
        except ValueError:
            due_date = first_date.replace(year=year, day=28)
        if due_date >= today:
            return due_date
    return None


def calculate_installment_info(first_insure_date_str, pay_period_years, due_date):
    first_date = datetime.datetime.strptime(first_insure_date_str, "%Y-%m-%d").date()
    current_installment = due_date.year - first_date.year + 1
    paid_installments = current_installment - 1
    return paid_installments, current_installment, pay_period_years


def get_policy_year(first_insure_date_str):
    """根据首保日期计算当前保单年度。first_insure_date 是第1保单年度。"""
    first_date = datetime.datetime.strptime(first_insure_date_str, "%Y-%m-%d").date()
    today = datetime.date.today()
    return today.year - first_date.year + 1


def encrypt_content(plaintext, password):
    """AES-GCM 加密文本"""
    salt = secrets.token_bytes(16)
    iv = secrets.token_bytes(12)
    key = hashlib.pbkdf2_hmac('sha256', password.encode('utf-8'), salt, 10000, dklen=32)
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    aesgcm = AESGCM(key)
    ciphertext = aesgcm.encrypt(iv, plaintext.encode('utf-8'), None)
    packed = salt + iv + ciphertext
    return base64.b64encode(packed).decode('ascii')


def get_policy_year_for_date(first_insure_date_str, target_year):
    """计算某个目标年份对应的保单年度"""
    first_date = datetime.datetime.strptime(first_insure_date_str, "%Y-%m-%d").date()
    return target_year - first_date.year + 1


def calc_premium_paid(first_insure_date_str, pay_period_years, premium, target_year, today):
    """
    计算截至目标年份的已交保费。
    - 过去/今年：按今天实际日期算（今天之前过了几个缴费日）
    - 未来年份：按截至该年年末算（保费 × 截至年末的期数）
    """
    first_date = datetime.datetime.strptime(first_insure_date_str, "%Y-%m-%d").date()
    premium = int(premium)

    if target_year <= today.year:
        # 过去或今年：按实际日期
        paid_periods = 0
        for period in range(pay_period_years):
            try:
                pay_date = first_date.replace(year=first_date.year + period)
            except ValueError:
                pay_date = first_date.replace(year=first_date.year + period, day=28)
            if pay_date <= today:
                paid_periods += 1
            else:
                break
    else:
        # 未来年份：截至该年年末
        year_end = datetime.date(target_year, 12, 31)
        paid_periods = 0
        for period in range(pay_period_years):
            try:
                pay_date = first_date.replace(year=first_date.year + period)
            except ValueError:
                pay_date = first_date.replace(year=first_date.year + period, day=28)
            if pay_date <= year_end:
                paid_periods += 1
            else:
                break

    return paid_periods * premium


def build_cash_value_html(data, cash_values):
    """构建现金价值汇总 HTML（支持年份筛选，含未来年份）"""
    if not cash_values:
        return ""

    today = datetime.date.today()
    current_year = today.year
    policies = data.get("policies", [])

    # 计算可选年份范围：从最早首保年到 Excel 中有数据的最远年份
    first_years = []
    max_data_year = current_year
    for p in policies:
        pid = p["policy_id"]
        if pid not in cash_values:
            continue
        fy = datetime.datetime.strptime(p["first_insure_date"], "%Y-%m-%d").year
        first_years.append(fy)
        # 找该保单数据的最大年度，换算成日历年
        if cash_values[pid]:
            max_policy_year = max(cash_values[pid].keys())
            max_calendar_year = fy + max_policy_year - 1
            max_data_year = max(max_data_year, max_calendar_year)

    if not first_years:
        return ""
    min_year = min(first_years)
    year_range = list(range(min_year, max_data_year + 1))

    # 预计算每个年份的数据
    all_year_data = {}
    for yr in year_range:
        total_cv = 0
        total_premium = 0
        details = []
        for p in policies:
            pid = p["policy_id"]
            if pid not in cash_values:
                continue
            first_date = datetime.datetime.strptime(p["first_insure_date"], "%Y-%m-%d").date()
            policy_year = yr - first_date.year + 1
            if policy_year < 1:
                continue
            cv = cash_values[pid].get(policy_year)
            if cv is None:
                continue
            # 特殊加成规则
            if pid == "POL004":
                if policy_year == 5:
                    cv += 10000
                elif policy_year > 5:
                    cv += 3000 * (policy_year - 5)
            elif pid == "POL006":
                if policy_year >= 6:
                    cv += (policy_year - 5) * 1000
            premium_paid = calc_premium_paid(
                p["first_insure_date"], p["pay_period_years"], p["premium"], yr, today
            )
            total_cv += cv
            total_premium += premium_paid
            details.append({
                "name": p["policy_name"],
                "user": p["user_name"],
                "policy_year": policy_year,
                "cv": cv,
                "premium_paid": premium_paid,
            })
        if details:
            all_year_data[yr] = {
                "total_cv": total_cv,
                "total_premium": total_premium,
                "details": details,
            }

    if not all_year_data:
        return ""

    # 年份快捷按钮：当年前后5年
    display_years = []
    for yr in range(current_year - 5, current_year + 6):
        if yr in all_year_data:
            display_years.append(yr)
    display_years.sort(reverse=True)

    year_tags = ""
    for yr in display_years:
        active = "active" if yr == current_year else ""
        year_tags += f"<span class='monthly-tag {active}' onclick='filterCashYear({yr}, event)'>{yr}</span>"

    # 自定义年份输入
    year_input = f'<input type="number" class="year-input" id="custom-year-input" min="{min_year}" max="{max_data_year}" placeholder="年份" onchange="filterCashYear(parseInt(this.value), event)">'

    # 为每个年份生成隐藏的数据 div
    year_divs = ""
    for yr, yd in all_year_data.items():
        display = "" if yr == current_year else "display:none;"
        rows = ""
        for d in yd["details"]:
            profit = d['cv'] - d['premium_paid']
            profit_class = "green" if profit >= 0 else "red"
            rows += f"<tr><td>{d['user']}</td><td>{d['name']}</td><td>{d['policy_year']}</td><td class='amount'>¥{d['cv']:,.0f}</td><td class='amount'>¥{d['premium_paid']:,}</td><td class='amount {profit_class}'>¥{profit:,.0f}</td></tr>"
        year_divs += f"""<div class="cash-year-block" data-cash-year="{yr}" style="{display}">
<div class="summary-cards">
<div class="summary-card"><div class="summary-label">总现金价值</div><div class="summary-value">¥{yd['total_cv']:,.0f}</div></div>
<div class="summary-card"><div class="summary-label">浮动盈亏</div><div class="summary-value {'green' if yd['total_cv'] - yd['total_premium'] >= 0 else 'red'}">¥{yd['total_cv'] - yd['total_premium']:,.0f}</div></div>
</div>
<table class="cash-table"><thead><tr>
<th>被保人</th><th>保单</th><th>保单年度</th><th>现金价值</th><th>已交保费</th><th>浮动盈亏</th>
</tr></thead><tbody>{rows}</tbody></table>
</div>"""

    html = f"""<div class="summary-section-title">保单现金价值</div>
<p style="font-size:10px;color:#999;margin:-6px 0 8px 2px;">POL004第5年+10000后每年+3000；POL006第6年起每年+1000</p>
<div class="monthly-detail"><span class="monthly-title">年份：</span>{year_input}<div class="monthly-tags">{year_tags}</div></div>
{year_divs}"""
    return html


def build_content_html(data, cash_values):
    today = datetime.date.today()
    policies = data.get("policies", [])

    policy_items = []
    for p in policies:
        next_due = calculate_next_due_date(p["first_insure_date"], p["pay_period_years"])
        if next_due is None:
            days_left_num = 999999
            paid = p["pay_period_years"]
            status = "✅ 已缴清"
            days_left = "-"
            status_class = "status-done"
            next_due_str = "-"
        else:
            paid, current, total = calculate_installment_info(
                p["first_insure_date"], p["pay_period_years"], next_due)
            days_left_num = (next_due - today).days
            days_left = f"{days_left_num} 天"
            next_due_str = str(next_due)
            if days_left_num <= 30:
                status = "🔴 即将缴费"
                status_class = "status-urgent"
            elif days_left_num <= 60:
                status = "🟡 临近缴费"
                status_class = "status-warning"
            else:
                status = "🟢 正常"
                status_class = "status-ok"
        policy_items.append({"p": p, "days_left_num": days_left_num, "paid": paid,
                             "status": status, "days_left": days_left,
                             "status_class": status_class, "next_due_str": next_due_str})

    policy_items.sort(key=lambda x: x["days_left_num"])

    grand_total = grand_paid = grand_remaining = 0
    for item in policy_items:
        p = item["p"]
        premium = int(p["premium"])
        total_periods = p["pay_period_years"]
        grand_total += premium * total_periods
        grand_paid += premium * item["paid"]
        grand_remaining += premium * (total_periods - item["paid"])

    current_year = today.year
    total_this_year = paid_this_year = unpaid_this_year = 0
    monthly_due = {}
    for item in policy_items:
        p = item["p"]
        next_due = calculate_next_due_date(p["first_insure_date"], p["pay_period_years"])
        if next_due is None:
            continue
        first_date = datetime.datetime.strptime(p["first_insure_date"], "%Y-%m-%d").date()
        try:
            this_year_due = first_date.replace(year=current_year)
        except ValueError:
            this_year_due = first_date.replace(year=current_year, day=28)
        end_year = first_date.year + p["pay_period_years"] - 1
        if first_date.year < current_year <= end_year:
            premium = int(p["premium"])
            total_this_year += premium
            if this_year_due < today:
                paid_this_year += premium
            else:
                unpaid_this_year += premium
                month = this_year_due.month
                monthly_due[month] = monthly_due.get(month, 0) + premium

    monthly_text = "<span class='monthly-tag active' onclick='filterMonth(0, event)'>全部</span>"
    for month in sorted(monthly_due.keys()):
        monthly_text += f"<span class='monthly-tag' onclick='filterMonth({month}, event)'>{month}月 ¥{monthly_due[month]:,}</span>"
    monthly_text = f"<div class='monthly-tags'>{monthly_text}</div>"

    policy_months = {}
    for item in policy_items:
        p = item["p"]
        next_due = calculate_next_due_date(p["first_insure_date"], p["pay_period_years"])
        if next_due is not None and next_due.year == current_year:
            policy_months[p["policy_id"]] = next_due.month
        else:
            policy_months[p["policy_id"]] = 0

    rows = ""
    for item in policy_items:
        p = item["p"]
        due_month = policy_months.get(p["policy_id"], 0)
        progress_pct = int(item['paid'] / p['pay_period_years'] * 100)
        row_class = ""
        if item["status_class"] == "status-urgent":
            row_class = "row-urgent"
        elif item["status_class"] == "status-warning":
            row_class = "row-warning"
        cloud_url = p.get("cloud_url", "")
        if cloud_url:
            view_btn = f'<a class="view-btn" href="{cloud_url}" target="_blank">查看</a>'
        else:
            view_btn = '<span class="no-file">暂无</span>'
        rows += f"""<tr data-month="{due_month}" class="{row_class}">
<td>{p['user_name']}</td><td><strong>{p['policy_name']}</strong></td><td>{p['company']}</td>
<td class="amount">¥{int(p['premium']):,}</td><td class="amount">¥{int(p['premium']) * item['paid']:,}</td><td class="amount">¥{int(p['premium']) * p['pay_period_years']:,}</td>
<td>{p['first_insure_date']}</td><td>{item['next_due_str']}</td>
<td><div class="progress-wrap"><div class="progress-bar" style="width:{progress_pct}%"></div><span class="progress-text">{item['paid']}/{p['pay_period_years']}期</span></div></td>
<td>{item['days_left']}</td><td class="{item['status_class']}">{item['status']}</td><td>{view_btn}</td>
</tr>"""

    # 现金价值汇总
    cash_value_html = build_cash_value_html(data, cash_values)

    content = f"""<div class="header">
<h1>📋 家庭保单管理</h1>
<p>更新时间：{today.strftime('%Y年%m月%d日')}</p>
</div>
<div class="tab-bar">
<span class="tab active" onclick="switchTab('overview', event)">资产概览</span>
<span class="tab" onclick="switchTab('payment', event)">缴费详情</span>
</div>
<div class="tab-content" id="tab-overview">
<div class="summary">
<div class="summary-section-title">累计保费</div>
<div class="summary-cards">
<div class="summary-card"><div class="summary-label">总保费</div><div class="summary-value">¥{grand_total:,}</div></div>
<div class="summary-card"><div class="summary-label">已交</div><div class="summary-value green">¥{grand_paid:,}</div></div>
<div class="summary-card"><div class="summary-label">剩余</div><div class="summary-value red">¥{grand_remaining:,}</div></div>
</div>
{cash_value_html}
</div>
</div>
<div class="tab-content" id="tab-payment" style="display:none;">
<div class="summary">
<div class="summary-section-title">{current_year}年保费</div>
<div class="summary-cards">
<div class="summary-card"><div class="summary-label">今年需交</div><div class="summary-value">¥{total_this_year:,}</div></div>
<div class="summary-card"><div class="summary-label">已交</div><div class="summary-value green">¥{paid_this_year:,}</div></div>
<div class="summary-card"><div class="summary-label">还需交</div><div class="summary-value red">¥{unpaid_this_year:,}</div></div>
</div>
<div class="monthly-detail"><span class="monthly-title">月度待缴：</span>{monthly_text}</div>
</div>
<div class="content">
<table><thead><tr>
<th>被保人</th><th>保单名称</th><th>保险公司</th><th>保费</th><th>已交保费</th><th>总保费</th><th>首保日期</th><th>下次续费</th><th>缴费进度</th><th>剩余天数</th><th>状态</th><th>保单详情</th>
</tr></thead><tbody>{rows}</tbody></table>
</div>
</div>
<div class="footer">自动生成 · 数据来源于保单管理系统</div>"""
    return content


def generate_html(data, password, cash_values):
    content_html = build_content_html(data, cash_values)
    encrypted_data = encrypt_content(content_html, password)

    css = load_template("style.css")
    js = load_template("app.js").replace("__ENCRYPTED_DATA__", encrypted_data)

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>保单续费详情</title>
<style>
{css}
</style>
</head>
<body>
<div id="login-screen" class="login-wrapper">
    <div class="login-box">
        <h2>🔒 保单详情</h2>
        <p>请输入访问密码查看保单信息</p>
        <input type="password" id="password-input" placeholder="请输入密码" onkeypress="if(event.key==='Enter')decrypt()">
        <button onclick="decrypt()">查看详情</button>
        <div id="error-msg" class="error-msg">密码错误，请重试</div>
    </div>
</div>
<div id="main-content" class="container" style="display:none;"></div>
<script>
{js}
</script>
</body>
</html>"""
    return html


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    data = load_config()
    cash_values = load_cash_values()
    html = generate_html(data, PAGE_PASSWORD, cash_values)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        f.write(html)


if __name__ == "__main__":
    main()
